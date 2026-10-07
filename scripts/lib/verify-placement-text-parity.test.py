import importlib.util
import tempfile
import struct
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE


SCRIPT_PATH = Path(__file__).resolve().parents[1] / "verify-placement-text-parity.py"
SPEC = importlib.util.spec_from_file_location("verify_placement_text_parity", SCRIPT_PATH)
PARITY = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(PARITY)


def font_named(name):
    encoded = name.encode("utf-16-be")
    names = struct.pack(">HHH", 0, 1, 18) + struct.pack(">HHHHHH", 3, 1, 0x409, 6, len(encoded), 0) + encoded
    return struct.pack(">IHHHH", 0x10000, 1, 0, 0, 0) + struct.pack(">4sIII", b"name", 0, 28, len(names)) + names


def reader_with_fonts(faces):
    fonts = {}
    operations = []
    for index, face in enumerate(faces):
        resource = f"/F{index}"
        fonts[resource] = {
            "/Subtype": "/TrueType", "/Encoding": "/WinAnsiEncoding",
            "/FontDescriptor": {"/FontFile2": SimpleNamespace(get_data=lambda name=face: font_named(name))},
        }
        operations.append(([resource, 11], b"Tf"))
    class Page(dict):
        def get_contents(self):
            return SimpleNamespace(operations=operations)
    return SimpleNamespace(pages=[Page({"/Resources": {"/Font": fonts}})])


class PlacementTextParityTests(unittest.TestCase):
    def test_require_is_active_under_optimized_python(self):
        with self.assertRaisesRegex(SystemExit, "required failure"):
            PARITY.require(False, "required failure")

    def test_forbidden_numeric_copy_uses_whole_number_boundaries(self):
        self.assertEqual(PARITY.find_forbidden_copy("30 contracts"), "30 contracts")
        self.assertEqual(PARITY.find_forbidden_copy("35 contracts"), "35 contracts")
        self.assertIsNone(PARITY.find_forbidden_copy("130 contracts and 235 contracts"))

    def test_normalization_joins_only_compounds_wrapped_after_hyphen_or_en_dash(self):
        self.assertEqual(PARITY.normalize_pdf_text("token-\nbased Oct 2022–\nPresent"), "token-based Oct 2022–Present")
        self.assertEqual(PARITY.normalize_pdf_text("Role —\nCompany"), "Role — Company")
        self.assertEqual(PARITY.normalize_pdf_text("token- based Oct 2022– Present"), "token- based Oct 2022– Present")

    def test_logical_links_merge_only_rectangles_owned_by_the_same_tagged_link(self):
        owner, separate_owner = {"/S": "/Link"}, {"/S": "/Link"}
        annotations = [
            {"/StructParent": key, "/A": {"/URI": "https://example.test/"}}
            for key in (1, 2, 3)
        ]
        reader = SimpleNamespace(
            trailer={"/Root": {"/StructTreeRoot": {"/ParentTree": {"/Kids": [
                {"/Nums": [1, owner, 2, owner, 3, separate_owner]},
            ]}}}},
            pages=[{"/Annots": annotations}],
        )
        self.assertEqual(PARITY.pdf_external_link_sequence(reader), ["https://example.test/", "https://example.test/"])
        annotations[1]["/A"]["/URI"] = "https://example.test/different"
        with self.assertRaisesRegex(SystemExit, "inconsistent destinations"):
            PARITY.pdf_external_link_sequence(reader)

    def test_unresolved_annotation_structure_parent_is_rejected(self):
        reader = SimpleNamespace(
            trailer={"/Root": {}},
            pages=[{"/Annots": [{"/StructParent": 7, "/A": {"/URI": "https://example.test/"}}]}],
        )
        with self.assertRaisesRegex(SystemExit, "unresolved structure parent"):
            PARITY.pdf_external_link_sequence(reader)

    def test_embedded_font_names_are_read_from_actual_truetype_bytes(self):
        self.assertEqual(PARITY.font_postscript_names(font_named("HPerkinsMarcellus-Regular")), {"HPerkinsMarcellus-Regular"})
        with self.assertRaisesRegex(SystemExit, "truncated"):
            PARITY.font_postscript_names(font_named("HPerkinsMarcellus-Regular")[:-1])

    def test_pdf_requires_all_six_theme_faces_and_embedded_searchable_fonts(self):
        reader = reader_with_fonts(sorted(PARITY.EXPECTED_FONT_FACES))
        self.assertEqual(PARITY.pdf_embedded_font_faces(reader), PARITY.EXPECTED_FONT_FACES)
        missing_face = reader_with_fonts(sorted(PARITY.EXPECTED_FONT_FACES)[:-1])
        with self.assertRaisesRegex(SystemExit, "all six"):
            PARITY.pdf_embedded_font_faces(missing_face)
        unapproved = reader_with_fonts(["Calibri"])
        with self.assertRaisesRegex(SystemExit, "unapproved embedded font"):
            PARITY.pdf_embedded_font_faces(unapproved)
        font = reader.pages[0]["/Resources"]["/Font"]["/F0"]
        font.pop("/Encoding")
        with self.assertRaisesRegex(SystemExit, "no readable text encoding"):
            PARITY.pdf_embedded_font_faces(reader)
        font["/Encoding"] = "/WinAnsiEncoding"
        reader.pages[0]["/Resources"]["/Font"]["/F0"]["/FontDescriptor"].clear()
        with self.assertRaisesRegex(SystemExit, "not embedded"):
            PARITY.pdf_embedded_font_faces(reader)

    def test_reads_nested_and_merged_table_cells_once_in_document_order(self):
        document = Document()
        document.add_paragraph("Letterhead")
        table = document.add_table(rows=1, cols=2)
        claim, metadata = table.rows[0].cells
        claim.paragraphs[0].text = "Claim"
        nested = claim.add_table(rows=1, cols=2)
        nested.cell(0, 0).merge(nested.cell(0, 1)).text = "One merged value"
        metadata.paragraphs[0].text = "Open"
        document.add_paragraph("End")

        paragraphs = list(PARITY.iter_document_paragraphs(document))
        self.assertEqual(
            [paragraph.text for paragraph in paragraphs if paragraph.text],
            ["Letterhead", "Claim", "One merged value", "Open", "End"],
        )
        self.assertEqual(len({id(paragraph._p) for paragraph in paragraphs}), len(paragraphs))

    def test_table_text_and_hyperlinks_preserve_visible_order_and_duplicates(self):
        document = Document()
        document.add_paragraph("Henry Perkins")
        table = document.add_table(rows=1, cols=2)
        for cell, text in zip(table.rows[0].cells, ("Evidence", "Released")):
            paragraph = cell.paragraphs[0]
            relationship_id = paragraph.part.relate_to(
                "https://example.test/evidence", RELATIONSHIP_TYPE.HYPERLINK, is_external=True,
            )
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), relationship_id)
            run = OxmlElement("w:r")
            text_node = OxmlElement("w:t")
            text_node.text = text
            run.append(text_node)
            hyperlink.append(run)
            paragraph._p.append(hyperlink)
        document.add_paragraph("End")

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "resume.docx"
            document.save(path)
            text, urls = PARITY.docx_text_and_urls(path)
        self.assertEqual(text, "Henry Perkins Evidence Released End")
        self.assertEqual(urls, ["https://example.test/evidence", "https://example.test/evidence"])

    def test_normalization_removes_temporary_pdf_after_write_failure(self):
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "resume.pdf"
            target.write_bytes(b"original")
            writer = mock.Mock()
            writer.write.side_effect = RuntimeError("write failed")
            with mock.patch.object(PARITY, "PdfReader"), mock.patch.object(PARITY, "PdfWriter", return_value=writer):
                with self.assertRaisesRegex(RuntimeError, "write failed"):
                    PARITY.normalize_word_pdf(target)
            self.assertEqual(list(Path(directory).glob("*.pdf")), [target])

    def test_explicit_pdf_path_is_used_for_normalization_and_verification(self):
        path = Path("selected.pdf")
        with mock.patch("sys.argv", ["verify-placement-text-parity.py", "--normalize-word-pdf", "--pdf", str(path)]):
            with mock.patch.object(PARITY, "normalize_word_pdf") as normalize:
                PARITY.main()
                normalize.assert_called_once_with(path)
        with mock.patch("sys.argv", ["verify-placement-text-parity.py", "--pdf", str(path)]):
            with mock.patch.object(PARITY, "verify") as verify:
                PARITY.main()
                verify.assert_called_once_with(path)


if __name__ == "__main__":
    unittest.main()
